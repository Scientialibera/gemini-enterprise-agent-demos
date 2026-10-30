"""Build the production IT requirements workbook: one tab per agent, plus the shared platform.

Scopes, roles, secrets and settings are read from deployment/registry.py, so the
workbook matches the code; rerun after changing an agent:

    uv run --group dev python docs/requirements/build_requirements.py

Writes docs/requirements/Prod_IT_Requirements.xlsx.
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "deployment"))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402
from registry import AGENTS, IDENTITY_OAUTH_SCOPES  # noqa: E402

OUT = Path(__file__).with_name("Prod_IT_Requirements.xlsx")
REDIRECT_URIS = (
    "https://vertexaisearch.cloud.google.com/static/oauth/oauth.html",
    "https://vertexaisearch.cloud.google.com/oauth-redirect",
)
BASELINE_ROLES = (
    "roles/aiplatform.expressUser",
    "roles/serviceusage.serviceUsageConsumer",
    "roles/parametermanager.parameterAccessor",
)
GUIDE = "docs/prod-runbook.md"
STATUSES = ("Not started", "Ticket raised", "In progress", "Done", "N/A", "Blocked")
COLUMNS = (
    ("#", 5),
    ("Area", 18),
    ("Requirement", 38),
    ("Exact value / specification", 60),
    ("Why it is needed", 46),
    ("Who provides it", 24),
    ("Automated by the repo?", 22),
    ("Status", 14),
    ("Ticket / owner / notes", 30),
)
FONT = "Arial"
HEADER_FILL = PatternFill("solid", start_color="1F3864")
AREA_FILL = PatternFill("solid", start_color="F2F2F2")
INPUT_FILL = PatternFill("solid", start_color="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Who provides what, in the client's terms.
GCP_ADMIN = "Cloud platform team"
IAM_ADMIN = "Cloud IAM / security"
GE_ADMIN = "Gemini Enterprise admin"
WS_ADMIN = "Google Workspace admin"
DATA_OWNER = "Data owner (BigQuery)"
ATL_ADMIN = "Atlassian admin"
DEPLOYER = "Deployment team (us)"
BUSINESS = "Business owner"

Row = tuple[str, str, str, str, str, str]  # area, requirement, value, why, who, automated


def _short(scope: str) -> str:
    return scope.rsplit("/", 1)[-1]


def oauth_rows(name: str) -> list[Row]:
    spec = AGENTS[name]
    scopes = (*IDENTITY_OAUTH_SCOPES, *spec.delegated_oauth_scopes)
    return [
        (
            "OAuth (user sign-in)",
            "Web OAuth client for this agent",
            "Google Auth Platform > Clients > Create > Web application. Name: <client name>. "
            f"Authorized redirect URIs (both): {REDIRECT_URIS[0]} and {REDIRECT_URIS[1]}. "
            "Client ID: <client ID>. Deliver the JSON (with the secret) through the approved "
            "secret channel, never email or chat.",
            "The agent acts as the signed-in user; Gemini Enterprise runs this OAuth flow the "
            "first time each user clicks Authorize. One client per agent: consent is cached "
            "per client.",
            GCP_ADMIN,
            "No (console only)",
        ),
        (
            "OAuth (user sign-in)",
            "Scopes on the OAuth consent screen",
            ", ".join(scopes) + f"  [short: {', '.join(_short(s) for s in scopes)}]",
            "Users are asked to consent to exactly these. Missing a scope breaks consent and "
            "cannot be fixed on an existing authorization without recreating it.",
            GCP_ADMIN,
            "No (console only)",
        ),
        (
            "OAuth (user sign-in)",
            "Workspace trust for the client (if user consent is restricted)",
            "Admin console > Security > API controls > App access control: mark the client "
            "ID as Trusted (or allow its scopes).",
            "Sensitive or restricted scopes (Drive, Docs, Sheets, cloud-platform) are often "
            "blocked for user consent in managed domains.",
            WS_ADMIN,
            "No (console only)",
        ),
        (
            "Gemini Enterprise",
            "Authorization resource",
            f"{spec.authorization_id} (created by the deployment from the client above)",
            "Links the agent to the OAuth client; the user's token reaches the agent through it.",
            DEPLOYER,
            "Yes (run.py deploy)",
        ),
    ]


def runtime_rows(name: str) -> list[Row]:
    spec = AGENTS[name]
    rows: list[Row] = [
        (
            "Runtime",
            "Agent Engine runtime with its own Agent Identity",
            f"Region of the deployment; package {spec.package_name}; identity type "
            "AGENT_IDENTITY (Google-managed, one per runtime).",
            "Hosts the code. Its Agent Identity is the agent's own principal, separate from "
            "users and from other agents.",
            DEPLOYER,
            "Yes (run.py deploy)",
        ),
        (
            "Runtime",
            "Runtime configuration parameter (prompt, model, limits)",
            f"Parameter Manager parameter {spec.config_parameter_id} (global). "
            + (
                "Stage prompts ship in the package (resources/prompts.yaml); any stage can be "
                "overridden through stage_instructions in this parameter. "
                if spec.is_workflow
                else "prompt.md is published to it on deploy. "
            )
            + "The runtime re-reads it every 30 seconds, so a prompt or model change needs "
            "no redeploy.",
            "Prompts and models are managed as configuration, versioned in Parameter Manager "
            "(and editable through the Prompt Manager agent).",
            DEPLOYER,
            "Yes (run.py deploy)",
        ),
        (
            "Gemini Enterprise",
            "Agent listing in the app",
            f"Display name: {spec.display_name}"
            + (
                "; plus a native workflow wrapper (backend listed as '… Backend')"
                if spec.is_workflow
                else ""
            ),
            "How users find and run the agent.",
            DEPLOYER,
            "Yes (run.py deploy)",
        ),
    ]
    for role in spec.agent_identity_project_roles:
        rows.append(
            (
                "IAM (agent identity)",
                f"Project role for this agent's identity: {role}",
                f"Grant {role} on the project to this runtime's Agent Identity only "
                "(principal printed by the deployment).",
                "Needed by this agent's own tools, beyond the shared baseline.",
                IAM_ADMIN,
                "Yes, if the deployer may set IAM",
            )
        )
    for grant in spec.agent_identity_secret_roles:
        rows.append(
            (
                "Secrets",
                f"Secret {grant.secret}",
                f"Secret Manager secret '{grant.secret}' with a value; "
                f"{', '.join(grant.roles)} granted on this secret only to this agent's identity.",
                "Read by the agent's tools at run time under its own identity; never in code or "
                "config files.",
                f"{IAM_ADMIN} + value owner",
                "Grant yes; value no",
            )
        )
    return rows


def user_access_row(name: str, extra: str = "") -> Row:
    return (
        "User access",
        "Who may use the agent",
        "Group: <Google group of intended users>, each with a Gemini Enterprise license; the "
        "agent is shared to that group (optional per-agent access groups are supported)." + extra,
        "Controls who sees and runs the agent.",
        BUSINESS + " / " + GE_ADMIN,
        "Partly (sharing yes, licensing no)",
    )


def acceptance_row(test: str) -> Row:
    return (
        "Acceptance",
        "End-to-end test",
        test,
        "Proves the setup works before users are switched over.",
        DEPLOYER,
        "Yes (run.py test)",
    )


def growth_rows() -> list[Row]:
    name = "growth_analysis_workflow"
    keys = {
        "growth-bls-api-key": "U.S. Bureau of Labor Statistics API key (register at bls.gov/developers)",
        "growth-fred-api-key": "FRED API key (fred.stlouisfed.org > My Account > API Keys)",
        "growth-nyt-api-key": "New York Times Article Search API key (developer.nytimes.com)",
        "growth-google-places-api-key": "Google Cloud API key restricted to Places API (New); "
        "the project that owns it must enable places.googleapis.com and have billing",
    }
    rows = runtime_rows(name)
    rows = [r if not r[1].startswith("Secret ") else _with_value_note(r, keys) for r in rows]
    return [
        *rows,
        *oauth_rows(name),
        (
            "Data (Google Drive)",
            "Store performance Google Sheet",
            "Spreadsheet ID: <spreadsheet ID>. The existing performance sheet with tabs "
            "'L6W SSS Summary (18 Mo.)' and 'L6W SSS Summary (15 Mo.)' (L6W date in E1, "
            "benchmark 'Simple Average' rows 3-6, owner tags on row 10, headers on row 11) and "
            "'ArcGIS (HH, At Scale GOV)' (headers on row 3). Headers that differ from the "
            "canonical names are mapped in configuration (column_aliases), not changed in "
            "the sheet.",
            "First stage reads it as the signed-in user: bottom 5 (18-month cohort), bottom 3 "
            "(15-month cohort) and top 5 stores, with benchmarks and owners for the TL;DR.",
            BUSINESS,
            "No",
        ),
        (
            "Data (Google Drive)",
            "Report folder (Shared Drive)",
            "Folder ID: <report folder ID>. A Shared Drive folder for the generated Google "
            "Docs; users of the agent need Contributor on it.",
            "The final stage creates the report Doc there as the signed-in user.",
            BUSINESS + " / " + WS_ADMIN,
            "No",
        ),
        (
            "Network",
            "Outbound HTTPS from Agent Engine",
            "api.bls.gov, api.stlouisfed.org, api.nytimes.com, archive-api.open-meteo.com, "
            "places.googleapis.com, and RSS: qsrmagazine.com, restaurantdive.com, "
            "fastcasual.com, *.eater.com (la, ny, philly).",
            "Research stages call these. Only relevant if egress is restricted (VPC Service "
            "Controls or a proxy).",
            GCP_ADMIN,
            "No",
        ),
        (
            "Quota",
            "Vertex AI model quota",
            "Gemini tokens per minute in the model location sized for the report stage, which "
            "sends the whole run's research (tens of thousands of tokens) in one request, "
            "times the number of runs at once. Rate-limited calls are retried for about a "
            "minute before a run fails.",
            "Without it, runs end without a report under load (429 RESOURCE_EXHAUSTED).",
            GCP_ADMIN,
            "No",
        ),
        user_access_row(
            name, " Users also need read on the sheet and Contributor on the report folder."
        ),
        (
            "Schedule (optional)",
            "Daily run",
            "Not enabled. A scheduled run has no signed-in user, so it would need a service "
            "account with access to the sheet and folder.",
            "The original deployment ran daily; decide whether that is still wanted.",
            BUSINESS,
            "No",
        ),
        acceptance_row(
            "'Run the pipeline' produces a Google Doc in the report folder, owned by the user."
        ),
    ]


def _with_value_note(row: Row, notes: dict[str, str]) -> Row:
    secret = row[1].removeprefix("Secret ")
    area, requirement, value, why, who, auto = row
    return (area, requirement, f"{value} Value: {notes.get(secret, 'see owner')}.", why, who, auto)


def vendor_rows() -> list[Row]:
    name = "vendor_match_agent"
    return [
        *runtime_rows(name),
        *oauth_rows(name),
        *[
            (
                "Data (Google Drive)",
                f"{label} folder (Shared Drive)",
                f"Folder ID: <{label.lower()} folder ID>. Shared Drive folder for {purpose}; "
                f"users need {access} on it.",
                "Every Drive call runs as the signed-in user.",
                BUSINESS + " / " + WS_ADMIN,
                "No",
            )
            for label, purpose, access in (
                ("Source", "incoming vendor lists", "Contributor"),
                ("Master dictionary", "the master vendor dictionary workbook", "Contributor"),
                ("Processed", "processed files and staged review batches", "Contributor"),
            )
        ],
        (
            "Data (Google Drive)",
            "Master vendor dictionary workbook",
            "The current master dictionary .xlsx in the master folder (file name unchanged).",
            "Lookups and approved mappings read and update it.",
            BUSINESS,
            "No",
        ),
        (
            "Quota",
            "Vertex AI model quota",
            "Gemini requests per minute in the deployment region sized for MAX_THREADS (default "
            "10) parallel resolver calls.",
            "Large batches resolve names in parallel with the model.",
            GCP_ADMIN,
            "No",
        ),
        user_access_row(name, " Users also need Contributor on the three folders."),
        acceptance_row("A dictionary lookup for a known vendor returns its master entry."),
    ]


def ticket_rows() -> list[Row]:
    name = "ticket_guidance_agent"
    return [
        *runtime_rows(name),
        *oauth_rows(name),
        (
            "Atlassian",
            "Dedicated Atlassian service account",
            "Account: <Atlassian service account email> on <Atlassian site URL>. An "
            "Atlassian account (not a person's) with: Jira Browse Projects, Create Issues, "
            "Modify Reporter, Transition Issues and Edit Issues in the target project; "
            "Confluence read on the knowledge base space.",
            "Jira is called as this account; the requesting user is set as reporter by email.",
            ATL_ADMIN,
            "No",
        ),
        (
            "Atlassian",
            "API token for that account",
            "id.atlassian.com > Security > API tokens (classic). Stored as the secret above; "
            "API base is the site URL (a scoped token needs api.atlassian.com/ex/jira/<cloudId>).",
            "Authenticates the Jira (and, in api mode, Confluence) calls.",
            ATL_ADMIN,
            "No",
        ),
        (
            "Atlassian",
            "Jira project and workflow",
            "Project key: <Jira project key>; issue types 'Service Request' and "
            "'Incident Request'; "
            "statuses 'Done' and 'In Progress' reachable from the initial status; optional "
            "parent epic.",
            "Tickets are raised there and moved to Done or In Progress by the conversation.",
            ATL_ADMIN,
            "Checked by run.py seed",
        ),
        (
            "Atlassian",
            "Knowledge base space",
            "Space key: <Confluence knowledge base space key>.",
            "Guidance is answered only from this space.",
            ATL_ADMIN,
            "No",
        ),
        (
            "Atlassian",
            "Atlassian OAuth 2.0 (3LO) app for the Confluence connector",
            "developer.atlassian.com > Create > OAuth 2.0 integration, resource-level. Callback: "
            "https://vertexaisearch.cloud.google.com/console/oauth/confluence_oauth.html. "
            "Confluence granular scopes: read:content-details, read:attachment, read:comment, "
            "read:page, read:whiteboard, read:space, read:group, read:user, "
            "read:content.metadata, read:configuration (each ':confluence'). Distribution: "
            "sharing, vendor Google, privacy policy https://policies.google.com.",
            "The Gemini Enterprise connector ingests Confluence with it.",
            ATL_ADMIN,
            "No (console only)",
        ),
        (
            "Gemini Enterprise",
            "Identity provider for access control",
            "Gemini Enterprise > Settings > Authentication: Google Identity for location global. "
            "Requires the Cloud project's organization to be a Workspace/Cloud Identity org.",
            "Ingested Confluence content is filtered by each user's Confluence permissions.",
            GE_ADMIN,
            "No (console only)",
        ),
        (
            "Gemini Enterprise",
            "Confluence Cloud connector (data ingestion)",
            "Data stores > Create > Confluence Cloud > Data ingestion > OAuth 2.0 client "
            "credentials. Instance URI = site URL; Instance ID = cloudId from "
            "<site>/_edge/tenant_info. Entities: Page, Blog, Comment. Full sync daily, "
            "incremental every 3 hours. Collection ID: <collection ID>.",
            "The agent searches its page, blog and comment data stores as the signed-in user.",
            GE_ADMIN,
            "No (console only)",
        ),
        user_access_row(name, " Users' emails must match their Atlassian accounts."),
        acceptance_row(
            "A known issue returns KB guidance and a ticket with the user as reporter; "
            "'resolved' moves it to Done."
        ),
    ]


def prompt_manager_rows() -> list[Row]:
    name = "parameter_updater"
    return [
        *runtime_rows(name),
        (
            "Governance",
            "Who may change prompts",
            "Restrict this agent to <prompt owners group>. Every change is shown as a diff "
            "and published only after confirmation; versions are never deleted or rewritten.",
            "It changes other agents' live behaviour without a deployment.",
            BUSINESS + " / " + GE_ADMIN,
            "Partly (sharing yes)",
        ),
        user_access_row(name, " Normally a small admin group only."),
        acceptance_row("Listing the agent configurations returns every code agent's parameter."),
    ]


def bigquery_mcp_rows(datasets: str) -> list[Row]:
    return [
        (
            "Gemini Enterprise",
            "BigQuery MCP connector (shared by both no-code agents)",
            "Custom MCP server https://bigquery.googleapis.com/mcp, OAuth (auth URI "
            "https://accounts.google.com/o/oauth2/v2/auth, extra params "
            "&access_type=offline&prompt=consent, token URI https://oauth2.googleapis.com/token, "
            "scope https://www.googleapis.com/auth/bigquery) with its own Web OAuth client "
            f"(both redirect URIs: {REDIRECT_URIS[0]}, {REDIRECT_URIS[1]}).",
            "The agent queries BigQuery through it as the signed-in user.",
            GE_ADMIN + " + " + GCP_ADMIN,
            "Scriptable (create_mcp_connector.py)",
        ),
        (
            "Org policy",
            "Allow custom MCP connectors",
            "constraints/discoveryengine.managed.disableCustomMcpServerConnector: enforce: false "
            "on the project.",
            "New organizations block custom MCP connectors by default.",
            IAM_ADMIN,
            "No",
        ),
        (
            "Gemini Enterprise",
            "Connect the connector's data store to the app",
            "Connected data stores > Add existing data stores > <connector>_mcp_data.",
            "Until then the agent designer does not offer the connector.",
            GE_ADMIN,
            "No (console only)",
        ),
        (
            "Gemini Enterprise",
            "Enable the connector's actions",
            "Data stores > the connector > Actions: execute_sql_readonly, list_dataset_ids, "
            "get_dataset_info, list_table_ids, get_table_info, get_query_results, get_job. "
            "Leave execute_sql off.",
            "With actions off the agent says the tool is unavailable.",
            GE_ADMIN,
            "No (console only)",
        ),
        (
            "User access",
            "BigQuery roles for every user",
            "roles/mcp.toolUser and roles/bigquery.jobUser on the project that runs the queries; "
            f"roles/bigquery.dataViewer on {datasets}.",
            "Queries run as the user; the agent cannot elevate.",
            IAM_ADMIN + " / " + DATA_OWNER,
            "No",
        ),
        (
            "Gemini Enterprise",
            "Publish the agent",
            "Agent designer > the agent > Create (first time) or Update. Connector changes need a "
            "recreate: deploy_no_code.py --recreate, then Create.",
            "The API saves drafts only; chat serves the last published version.",
            GE_ADMIN + " / " + DEPLOYER,
            "Draft yes; publish no",
        ),
    ]


def kitchen_rows() -> list[Row]:
    return [
        (
            "Gemini Enterprise",
            "No-code agent",
            "agents/no_code/kitchen_sequencer (instruction, model, starter prompts, connector).",
            "Kitchen compliance insights from curated tables, replacing the Looker dashboards.",
            DEPLOYER,
            "Yes (run.py deploy), then one Create/Update click",
        ),
        (
            "Data (BigQuery)",
            "Curated dataset and daily refresh",
            "Dataset kitchen_compliance (US) with 4 tables rebuilt daily by BigQuery scheduled "
            "queries (infrastructure/sql/kitchen_compliance), run by service account "
            "kitchen-compliance-refresh@<project>.iam.gserviceaccount.com "
            "(bigquery.dataEditor on the dataset, bigquery.jobUser on the project).",
            "The agent answers only from these tables.",
            DEPLOYER + " (Terraform)",
            "Yes (Terraform)",
        ),
        (
            "Data (BigQuery)",
            "Source data read for the refresh account",
            "roles/bigquery.dataViewer on <source project>.<source dataset> (the order and "
            "item compliance tables) for "
            "kitchen-compliance-refresh@<project>.iam.gserviceaccount.com. APIs: "
            "bigquerydatatransfer.googleapis.com.",
            "The scheduled queries read the order and item compliance tables.",
            DATA_OWNER,
            "No",
        ),
        *bigquery_mcp_rows("the kitchen_compliance dataset"),
        user_access_row("kitchen_sequencer"),
        acceptance_row(
            "'What is the Kitchen Compliance Scorecard for this week?' matches the Looker figures for "
            "the same week."
        ),
    ]


def recipe_rows() -> list[Row]:
    return [
        (
            "Gemini Enterprise",
            "No-code agent",
            "agents/no_code/recipe_editorial_agent (instruction, model, connectors).",
            "Editorial QA of recipes with an emailed HTML report.",
            DEPLOYER,
            "Yes (run.py deploy), then one Create/Update click",
        ),
        (
            "Data (BigQuery)",
            "Recipe datasets",
            "Product dataset: <project>.<product recipe dataset> with tables (or views) "
            "culinary_products, culinary_product_preparation_steps, "
            "culinary_product_ingredients, culinary_ingredient_specifications, menu_cycles, "
            "menu_items; legacy dataset: <project>.<legacy recipe dataset> with "
            "legacy_recipes, legacy_recipe_steps, legacy_recipe_ingredients, "
            "legacy_cycle_recipe_slots.",
            "The prompt reads recipes, steps, ingredients and cycles from them.",
            DATA_OWNER,
            "No",
        ),
        *bigquery_mcp_rows("both recipe datasets"),
        (
            "Gemini Enterprise",
            "Gmail connector (optional, to email the report)",
            "Data stores > Create > Gmail (federated, with actions); Workspace admin consent "
            "for the send action.",
            "Without it the report is shown in chat instead of emailed.",
            GE_ADMIN + " + " + WS_ADMIN,
            "No (console only)",
        ),
        user_access_row("recipe_editorial_agent"),
        acceptance_row(
            "A recipe SKU returns the QA report; with Gmail, it is emailed to the recipient."
        ),
    ]


def platform_rows() -> list[Row]:
    apis = (
        "aiplatform, discoveryengine, parametermanager, secretmanager, serviceusage, storage, "
        "logging, monitoring, cloudtrace, telemetry, docs, drive, sheets, bigquery, "
        "bigquerydatatransfer (each .googleapis.com); places.googleapis.com in the project that "
        "owns the Places key."
    )
    return [
        (
            "Project",
            "Google Cloud project",
            "Project ID: <project ID>. A dedicated project in the client's organization, "
            "with billing.",
            "Hosts every runtime, secret, parameter and dataset.",
            GCP_ADMIN,
            "No",
        ),
        (
            "Project",
            "Required APIs",
            apis,
            "The agents and the deployment use them.",
            GCP_ADMIN,
            "Yes (Terraform or run.py)",
        ),
        (
            "Gemini Enterprise",
            "Gemini Enterprise app and licences",
            "App (engine) ID: <app ID>, in location global (allowed by org policy); licences "
            "for every user.",
            "All agents are published into this app.",
            GE_ADMIN,
            "No (console only)",
        ),
        (
            "Gemini Enterprise",
            "Identity provider",
            "Settings > Authentication: Google Identity for global (Workspace or Cloud Identity org).",
            "Needed by ingestion connectors (Confluence) to enforce source permissions.",
            GE_ADMIN,
            "No (console only)",
        ),
        (
            "OAuth",
            "OAuth consent screen",
            "Google Auth Platform: Internal user type; app name <app name> and support "
            "email <support email>; Data Access lists every scope on the agent tabs.",
            "Every delegated agent and the BigQuery MCP connector use it.",
            GCP_ADMIN,
            "No (console only)",
        ),
        (
            "IAM (all agents)",
            "Baseline roles for every Agent Identity",
            ", ".join(BASELINE_ROLES)
            + " on the project, granted to the principal set of all Agent "
            "Identities in the project; roles/storage.objectViewer on the staging bucket only.",
            "Every runtime reads its own parameter and model with these; without them it deploys "
            "and then fails on the first request.",
            IAM_ADMIN,
            "Yes (Terraform)",
        ),
        (
            "Storage",
            "Staging bucket",
            "gs://<staging bucket> in <region> for runtime packages.",
            "Agent Engine deploys code from it.",
            GCP_ADMIN,
            "Yes (Terraform or run.py)",
        ),
        (
            "Secrets",
            "Secret Manager secrets",
            "Created empty by Terraform (managed_secrets); values supplied through the approved "
            "secret process. See the Secrets tab.",
            "API keys and the Jira token; never in code.",
            IAM_ADMIN,
            "Create yes; values no",
        ),
        (
            "Deployment access",
            "Deployer identity (person or CI service account)",
            "Principal: <deployer user or CI service account>. On the project: "
            "roles/aiplatform.user, roles/discoveryengine.admin, "
            "roles/parametermanager.parameterAdmin, roles/secretmanager.admin, "
            "roles/storage.admin (staging bucket), roles/bigquery.dataEditor + jobUser (seeding), "
            "roles/serviceusage.serviceUsageConsumer; roles/resourcemanager.projectIamAdmin only "
            "if it should grant the per-agent roles itself.",
            "Runs `deployment/run.py` (deploy, register, publish prompts, seed, test).",
            IAM_ADMIN,
            "No",
        ),
        (
            "Observability",
            "Logging and monitoring",
            "Cloud Logging and Monitoring on the project; optional log sink and alerts (Terraform "
            "observability module).",
            "Operating and auditing the agents.",
            GCP_ADMIN,
            "Yes (Terraform)",
        ),
        (
            "Org policy",
            "Custom MCP connectors allowed",
            "constraints/discoveryengine.managed.disableCustomMcpServerConnector enforce: false.",
            "Needed by the BigQuery MCP connector (no-code agents).",
            IAM_ADMIN,
            "No",
        ),
        (
            "Org policy",
            "Allowed locations",
            "global (Gemini Enterprise) and the Agent Engine region <region>.",
            "Resource creation fails in a disallowed location.",
            IAM_ADMIN,
            "No",
        ),
    ]


SHEETS: dict[str, tuple[str, list[Row]]] = {
    "Shared Platform": ("Project-wide prerequisites every agent depends on.", platform_rows()),
    "Growth Analysis": (
        "ADK workflow: store performance diagnostic report (Google Doc).",
        growth_rows(),
    ),
    "Vendor Match": (
        "ADK agent: vendor name standardization against the Drive master dictionary.",
        vendor_rows(),
    ),
    "Ticket Guidance": ("ADK agent: Confluence KB guidance and Jira tickets.", ticket_rows()),
    "Prompt Manager": (
        "ADK agent: view and change code agents' prompts and models in Parameter Manager.",
        prompt_manager_rows(),
    ),
    "Kitchen Sequencer": (
        "No-code agent: kitchen compliance from curated BigQuery tables.",
        kitchen_rows(),
    ),
    "Recipe Editorial": (
        "No-code agent: recipe editorial QA from BigQuery, emailed report.",
        recipe_rows(),
    ),
}


def style_header(ws, row: int) -> None:
    for col, (title, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = width


def write_requirements(ws, title: str, summary: str, rows: list[Row]) -> None:
    ws["A1"] = title
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = summary
    ws["A2"].font = Font(name=FONT, italic=True)
    ws["A3"] = (
        "Values in <angle brackets> are placeholders for the client's IT team to decide and "
        "fill in (those cells are yellow, like Status and notes); everything else is fixed by "
        f"Google or by the repository. Details: {GUIDE}."
    )
    ws["A3"].font = Font(name=FONT, size=9, color="595959")
    header = 5
    style_header(ws, header)
    validation = DataValidation(
        type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True
    )
    ws.add_data_validation(validation)
    for i, row in enumerate(rows, start=1):
        r = header + i
        values = (i, *row, "Not started", "")
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER
            if col == 2:
                cell.fill = AREA_FILL
            if col in (8, 9) or (col == 4 and "<" in str(value)):
                cell.fill = INPUT_FILL
        validation.add(f"H{r}")
    ws.freeze_panes = ws.cell(row=header + 1, column=3)
    ws.auto_filter.ref = f"A{header}:{get_column_letter(len(COLUMNS))}{header + len(rows)}"


def write_secrets(ws) -> None:
    ws["A1"] = "Secrets register"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = "Every secret the agents read. Values never go in code, config files, email or chat."
    ws["A2"].font = Font(name=FONT, italic=True)
    headers = (
        ("Secret ID", 34),
        ("Used by", 24),
        ("Value / source", 60),
        ("Read by", 40),
        ("Rotation owner", 24),
        ("Status", 14),
    )
    for col, (title, width) in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col, value=title)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = width
    rows = []
    for spec in AGENTS.values():
        for grant in spec.agent_identity_secret_roles:
            rows.append(
                (
                    grant.secret,
                    spec.display_name,
                    "See the agent's tab",
                    f"{spec.display_name} Agent Identity ({', '.join(grant.roles)}, this secret only)",
                    "Value owner",
                )
            )
        if spec.uses_delegated_auth:
            rows.append(
                (
                    f"{spec.package_name}-oauth-client-secret",
                    spec.display_name,
                    "The agent's Google OAuth client secret (imported on first registration)",
                    "Gemini Enterprise authorization (deployment)",
                    GCP_ADMIN,
                )
            )
    rows.append(
        (
            "(connector-held) BigQuery MCP OAuth client secret",
            "Kitchen Sequencer, Recipe Editorial",
            "Entered once into the connector; held by Gemini Enterprise",
            "Gemini Enterprise",
            GCP_ADMIN,
        )
    )
    rows.append(
        (
            "(connector-held) Atlassian OAuth app secret",
            "Ticket Guidance",
            "Entered once into the Confluence connector; held by Gemini Enterprise",
            "Gemini Enterprise",
            ATL_ADMIN,
        )
    )
    validation = DataValidation(
        type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True
    )
    ws.add_data_validation(validation)
    for i, row in enumerate(rows, start=5):
        for col, value in enumerate((*row, "Not started"), start=1):
            cell = ws.cell(row=i, column=col, value=value)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER
        ws.cell(row=i, column=6).fill = INPUT_FILL
        validation.add(f"F{i}")
    ws.freeze_panes = "B5"


def write_overview(ws) -> None:
    ws["A1"] = "Gemini Enterprise agents: production IT requirements"
    ws["A1"].font = Font(name=FONT, bold=True, size=16)
    lines = (
        "One tab per agent, plus the shared platform every agent needs and a secrets register.",
        "Each row is one thing to request or provide: what, the exact value, why, who, and "
        "whether the repository automates it. Track progress in the yellow Status and "
        "Ticket columns; the table below totals them.",
        "Agent configuration (prompt, model, limits) for code agents lives in Parameter Manager "
        "and is published from the repository on each deployment.",
        f"Generated from deployment/registry.py by docs/requirements/build_requirements.py; "
        f"the step-by-step runbook is {GUIDE}.",
    )
    for i, line in enumerate(lines, start=3):
        cell = ws.cell(row=i, column=1, value=line)
        cell.font = Font(name=FONT, size=10)
        cell.alignment = Alignment(wrap_text=True)
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=7)
        ws.row_dimensions[i].height = 30
    header = 9
    titles = ("Tab", "What it covers", "Requirements", "Done", "N/A", "Open", "% complete")
    widths = (22, 70, 14, 10, 10, 10, 12)
    for col, (title, width) in enumerate(zip(titles, widths, strict=True), start=1):
        cell = ws.cell(row=header, column=col, value=title)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = width
    for i, (tab, (summary, _)) in enumerate(SHEETS.items(), start=header + 1):
        ref = f"'{tab}'!H:H"
        values = (
            tab,
            summary,
            f"=COUNTA({ref})-1",
            f'=COUNTIF({ref},"Done")',
            f'=COUNTIF({ref},"N/A")',
            f"=C{i}-D{i}-E{i}",
            f'=IF(C{i}-E{i}=0,"-",D{i}/(C{i}-E{i}))',
        )
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=i, column=col, value=value)
            cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=i, column=7).number_format = "0%"
    total = header + len(SHEETS) + 1
    ws.cell(row=total, column=1, value="Total").font = Font(name=FONT, bold=True)
    for col in range(3, 7):
        letter = get_column_letter(col)
        cell = ws.cell(
            row=total, column=col, value=f"=SUM({letter}{header + 1}:{letter}{total - 1})"
        )
        cell.font = Font(name=FONT, bold=True)
    cell = ws.cell(
        row=total, column=7, value=f'=IF(C{total}-E{total}=0,"-",D{total}/(C{total}-E{total}))'
    )
    cell.font = Font(name=FONT, bold=True)
    cell.number_format = "0%"
    legend = total + 2
    ws.cell(row=legend, column=1, value="Status values").font = Font(name=FONT, bold=True)
    ws.cell(row=legend, column=2, value=", ".join(STATUSES)).font = Font(name=FONT, size=10)
    ws.cell(row=legend + 1, column=1, value="Who provides it").font = Font(name=FONT, bold=True)
    ws.cell(
        row=legend + 1,
        column=2,
        value=", ".join(
            (GCP_ADMIN, IAM_ADMIN, GE_ADMIN, WS_ADMIN, DATA_OWNER, ATL_ADMIN, BUSINESS, DEPLOYER)
        ),
    ).font = Font(name=FONT, size=10)


def main() -> None:
    wb = Workbook()
    write_overview(wb.active)
    wb.active.title = "Overview"
    for tab, (summary, rows) in SHEETS.items():
        write_requirements(wb.create_sheet(tab), tab, summary, rows)
    write_secrets(wb.create_sheet("Secrets"))
    wb.save(OUT)
    print(f"WROTE={OUT.relative_to(ROOT)} tabs={len(wb.sheetnames)}")


if __name__ == "__main__":
    main()
