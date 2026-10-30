"""Recipe steps layout helpers."""

from __future__ import annotations

import functools
import math
from typing import Any, NamedTuple

from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from ..errors import ContentTooLong
from .content import cooking_tip_for_steps, split_instructions
from .drawing import (
    CHECKLIST_INDENT,
    add_box,
    add_checkbox,
    add_checklist,
    add_image,
    add_ink_image,
    add_rule,
    add_soft_shadow,
    add_text,
    add_vrule,
    clean,
    limit_text,
    no_shadow,
    rgb,
    round_picture,
    wrapped_line_count,
)
from .footer import add_footer
from .ingredients import add_ingredient_strip, ingredient_strip_height
from .theme import (
    BANNER_CHAR_WIDTH_RATIO,
    BANNER_FONT_SIZE,
    BANNER_H,
    BANNER_TEXT_W,
    BANNER_TEXT_X,
    BODY_FONT,
    BOLD_WIDTH_FACTOR,
    BULLET_FONT_SIZE,
    BULLET_LINE_HEIGHT,
    BULLET_MIN_HEIGHT,
    BULLET_ROW_PADDING,
    BULLET_WRAP_ALLOWANCE,
    CHAR_WIDTH_RATIO,
    CLOSING_BOTTOM,
    CLOSING_GAP,
    HEAD_CHAR_WIDTH_RATIO,
    HEAD_FONT,
    INLINE_COLUMN_GAP,
    INLINE_FONT_SIZE,
    LINE_HEIGHT_RATIO,
    POINTS_PER_INCH,
    STACK_BLOCK_MAX_SCALE,
    STACK_COLUMNS,
    STACK_GAP,
    STACK_PHOTO_ASPECT,
    STACK_TITLE_BOX_HEIGHT,
    STACK_TITLE_FONT_SIZE,
    STEP_BADGE_D,
    STEP_BLOCK_GAP,
    STEP_BLOCK_MAX_SCALE,
    STEP_BLOCK_MIN_HEIGHT,
    STEP_IMAGE_MAX_PORTRAIT,
    STEP_IMAGE_MIN_ASPECT,
    STEP_IMAGE_WIDTH_FRACTION,
    STEP_TITLE_BOX_HEIGHT,
    STEP_TITLE_FONT_SIZE,
    STEP_TITLE_ONE_LINE_H,
    STEPS_PAGE_BOTTOM,
    TIP_CHARS_PER_LINE,
    TIP_LINE_HEIGHT,
    TIP_VERTICAL_PADDING,
    VARIATIONS_INLINE_MAX_H,
    VARIATIONS_PANEL_H,
    VARIATIONS_PANEL_OFFSET,
    VARIATIONS_SKETCH_X,
    VARIATIONS_TAB_FILL,
    VARIATIONS_TAB_H,
    VARIATIONS_TAB_W,
    WORDS_PER_BULLET_LINE,
    C,
)

LAYOUT_SIDE = "side"
LAYOUT_STACKED = "stacked"
# Two tips of 150 characters can share a page now pages are packed by height.
TIP_MAX_CHARS = 320
CONTENT_X = 0.34
CONTENT_W = 9.22
# A block wider than this is a single step spanning the page.
WIDE_BLOCK_W = 6.0
# The band and notes that close a page whose steps continue overleaf.
CONTINUED_BAND_H = 0.56
CONTINUED_BAND_BOTTOM = 12.54
CONTINUED_MIN_SPARE = 0.9
NOTES_MIN_H = 1.6
NOTES_LINE_GAP = 0.42
# (characters of the longest step on the page, photograph share) from short
# to long; longer pages use STEP_IMAGE_WIDTH_FRACTION.
PHOTO_FRACTION_BY_TEXT = ((170, 0.56), (300, 0.51))
STEP_IMAGE_WIDE_ASPECT = 0.62
# A step in a row across the page: its photograph's share of the row, how tall
# that landscape photograph may grow, and the most steps set this way.
WIDE_PHOTO_FRACTION = 0.48
STEP_IMAGE_WIDE_MAX = 0.74
WIDE_MAX_STEPS = 3


def tip_box_height(text: str) -> float:
    """Height of the tip banner for this text, or 0 when the page has no tip."""
    body = clean(text).strip()
    if not body:
        return 0.0
    tip = f"Cooking Tip: {limit_text(body, TIP_MAX_CHARS)}"
    lines = max(1, math.ceil(len(tip) / TIP_CHARS_PER_LINE))
    return lines * TIP_LINE_HEIGHT + TIP_VERTICAL_PADDING


def _tip_rule_y(text: str) -> float:
    box_h = tip_box_height(text)
    return 0.30 + (box_h + 0.10 if box_h else TIP_VERTICAL_PADDING)


def steps_top(text: str) -> float:
    """Where a page's steps begin, below its tip banner and rule."""
    return _tip_rule_y(text) + 0.18


def add_top_tip(slide, text, *, split=True):
    """Draw the tip banner, or nothing when this page has no tip.

    A page without a tip returns the line the steps would start under anyway,
    so a missing tip closes the gap rather than leaving an empty yellow bar.
    """
    rule_y = _tip_rule_y(text)
    box_h = tip_box_height(text)
    if box_h:
        box = add_box(slide, 0.34, 0.30, 9.22, box_h, C["yellow2"], C["yellow2"], radius=True)
        add_soft_shadow(box)
        # Only the label is set as a heading. A whole sentence in bold serif
        # shouted over the steps; the reference card sets the tip itself in
        # the body face, the way a margin note reads.
        label = add_text(
            slide,
            "Cooking Tip:  ",
            0.53,
            0.30,
            8.85,
            box_h,
            font_face=HEAD_FONT,
            font_size=15.5,
            color=C["dark_blue"],
            bold=True,
            valign="middle",
            fit=False,
        )
        run = label.text_frame.paragraphs[0].add_run()
        run.text = limit_text(clean(text).strip(), TIP_MAX_CHARS)
        run.font.name = BODY_FONT
        run.font.size = Pt(12)
        run.font.color.rgb = rgb(C["dark_blue"])

    # Broken over the column gutter only when the steps below sit in columns.
    if split:
        add_rule(slide, 0.34, rule_y, 4.55, C["line"], 0.8)
        add_rule(slide, 5.12, rule_y, 4.44, C["line"], 0.8)
    else:
        add_rule(slide, 0.34, rule_y, 9.22, C["line"], 0.8)
    return rule_y


def _bullet_heights(
    bullets: list[str], width: float, minimum: float, padding: float = BULLET_ROW_PADDING
) -> list[float]:
    """Height each instruction needs at the bullet size, wrapped to ``width``.

    Measured a little narrower than the box, because PowerPoint breaks lines
    slightly earlier than the mean character width predicts: a bullet counted
    as three lines set as four and ran into the one below it.
    """
    measure_w = max(0.5, width - BULLET_WRAP_ALLOWANCE)
    return [
        max(
            minimum,
            wrapped_line_count(bullet, measure_w, BULLET_FONT_SIZE, CHAR_WIDTH_RATIO)
            * BULLET_FONT_SIZE
            * LINE_HEIGHT_RATIO
            / POINTS_PER_INCH
            + padding,
        )
        for bullet in bullets
    ]


def _title_height(step: dict[str, Any], width: float) -> float:
    """Height a step title takes: one line when it fits, two otherwise.

    Measured with the variation marker added, so a title that gains " *" on
    the page cannot wrap where the measurement said it would not.
    """
    title = clean(step.get("title"), "Step") + " *"
    lines = wrapped_line_count(
        title, width - STEP_BADGE_D - 0.10, STEP_TITLE_FONT_SIZE, HEAD_CHAR_WIDTH_RATIO
    )
    return STEP_TITLE_ONE_LINE_H if lines == 1 else STEP_TITLE_BOX_HEIGHT


def step_block_height(
    step: dict[str, Any],
    width: float,
    fraction: float = STEP_IMAGE_WIDTH_FRACTION,
    aspect: float | None = None,
) -> float:
    """Height this step needs: its title, its bullets and its photograph.

    Steps differ in length, so each block is measured rather than given a
    share of a fixed grid. The photograph sets a floor, because a block
    shorter than its own image would crop it. ``aspect`` is the photograph's
    height over its width; by default the smallest the block's shape allows.
    """
    bullets = split_instructions(step)
    image_w = width * fraction
    text_w = width - image_w - 0.16 - 0.30

    # Measured exactly as the block draws it. A separate rough estimate here
    # counted each line at twice its real height, so four steps that fitted
    # one page were split two and two across pages of white.
    text_h = sum(_bullet_heights(bullets, text_w, BULLET_MIN_HEIGHT))

    # The photograph keeps a portrait proportion of its own. A step with one or
    # two instructions would otherwise get a block barely taller than its text,
    # and a noticeably smaller picture than the steps beside it, which reads as
    # a mistake rather than as a short step.
    if aspect is None:
        aspect = STEP_IMAGE_WIDE_ASPECT if width > WIDE_BLOCK_W else STEP_IMAGE_MIN_ASPECT
    content_h = _title_height(step, width) + 0.10 + max(text_h, image_w * aspect)
    return max(STEP_BLOCK_MIN_HEIGHT, content_h + 0.20)


def _tallest_aspect(width: float) -> float:
    """The tallest a photograph may be drawn, over its width, at this block width."""
    return STEP_IMAGE_WIDE_MAX if width > WIDE_BLOCK_W else STEP_IMAGE_MAX_PORTRAIT


def stacked_column_width() -> float:
    return (CONTENT_W - (STACK_COLUMNS - 1) * STACK_GAP) / STACK_COLUMNS


def stacked_block_height(step: dict[str, Any], width: float) -> float:
    """Height of a compact block: title, then instructions, then the photograph."""
    bullets = split_instructions(step)
    text_h = sum(_bullet_heights(bullets, width - 0.26, 0.24))
    return STACK_TITLE_BOX_HEIGHT + 0.08 + text_h + 0.12 + width * STACK_PHOTO_ASPECT + 0.20


def _add_step_heading(slide, idx, step, x, y, w, *, box_h, font_size, badge_d):
    """Number disc and title. The disc lets the eye find step 3 at a glance."""
    badge = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x), Inches(y + 0.03), Inches(badge_d), Inches(badge_d)
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = rgb(C["dark_blue"])
    badge.line.fill.background()
    no_shadow(badge)
    add_text(
        slide,
        str(idx + 1),
        x,
        y + 0.03,
        badge_d,
        badge_d,
        font_size=13 if badge_d >= STEP_BADGE_D else 11.5,
        color=C["white"],
        bold=True,
        align="center",
        fit=False,
    )
    # Two lines, so a longer step name sets at the same size as a short one.
    indent = badge_d + 0.10
    title = clean(step.get("title"), "Step")
    # A one-line title is centred on its number disc; set at the top of a
    # two-line box it rode visibly above the disc's centre.
    one_line = wrapped_line_count(title, w - indent, font_size, HEAD_CHAR_WIDTH_RATIO) == 1
    add_text(
        slide,
        title,
        x + indent,
        y + 0.03 if one_line else y,
        w - indent,
        badge_d if one_line else box_h,
        font_face=HEAD_FONT,
        font_size=font_size,
        color=C["dark_blue"],
        bold=True,
        valign="middle" if one_line else "top",
    )


def _add_step_photo(slide, step, x, y, w, h):
    photo = add_image(
        slide,
        step.get("image_path") or step.get("imagePath"),
        x,
        y,
        w,
        h,
        crop=True,
        placeholder="STEP IMAGE",
    )
    if photo is not None:
        round_picture(photo)
        add_soft_shadow(photo, alpha=24)


def _add_bullets(slide, idx, bullets, heights, x, top, width, available_h):
    """Draw the checklist, or ask for a shorter step when it cannot fit."""
    if sum(heights) > available_h + 0.01:
        # The model can fix this, so say which step and by how much rather than
        # failing with a generic message it cannot act on.
        overflow = sum(heights) - available_h
        words_over = max(1, round(overflow / BULLET_LINE_HEIGHT * WORDS_PER_BULLET_LINE))
        raise ContentTooLong(
            field=f"steps[{idx}].body",
            detail=(
                f"Step {idx + 1} needs {sum(heights):.2f}in of text area but has "
                f"{available_h:.2f}in."
            ),
            suggestion=(
                f"Remove about {words_over} words from step {idx + 1}, or split it into two steps."
            ),
        )
    add_checklist(
        slide,
        bullets,
        x,
        top - 0.02,
        width + CHECKLIST_INDENT,
        available_h,
        font_size=BULLET_FONT_SIZE,
        space_after=BULLET_ROW_PADDING * POINTS_PER_INCH,
    )


def add_step_block(slide, step, idx, x, y, w, h, *, fraction=STEP_IMAGE_WIDTH_FRACTION):
    """Side-by-side block: the checklist on the left, the photograph beside it."""
    _add_step_heading(
        slide,
        idx,
        step,
        x,
        y,
        w,
        box_h=STEP_TITLE_BOX_HEIGHT,
        font_size=STEP_TITLE_FONT_SIZE,
        badge_d=STEP_BADGE_D,
    )

    # Always reserve a standardized image area for every step, even when an image
    # has not yet been generated. This is deliberate: the image agent can fill the
    # placeholder later without changing layout geometry.
    # The reference cards set the step photograph in portrait, filling the
    # column beside the text rather than sitting in a fixed square.
    img_w = w * fraction
    img_x = x + w - img_w
    bullet_top = y + _title_height(step, w) + 0.06
    img_y = bullet_top + 0.04
    # The photograph grows with a taller block, but only to a portrait
    # proportion: beyond it the crop turns a square image into a sliver.
    # A step alone across the full width takes a landscape photograph: at that
    # width a portrait one filled most of the page.
    img_h = min(h - (img_y - y) - 0.12, img_w * _tallest_aspect(w))
    _add_step_photo(slide, step, img_x, img_y, img_w, img_h)

    bullets = split_instructions(step)
    bullet_w = w - img_w - 0.16 - 0.30
    _add_bullets(
        slide,
        idx,
        bullets,
        _bullet_heights(bullets, bullet_w, BULLET_MIN_HEIGHT),
        x,
        bullet_top,
        bullet_w,
        h - (bullet_top - y) - 0.16,
    )
    add_rule(slide, x, y + h - 0.05, w, C["line"], 0.55)


def add_stacked_block(slide, step, idx, x, y, w, h):
    """Compact block: title, checklist, then the photograph across the column.

    The photograph sits on the block's floor, so in a row of three the
    pictures line up whatever the length of the text above them.
    """
    _add_step_heading(
        slide,
        idx,
        step,
        x,
        y,
        w,
        box_h=STACK_TITLE_BOX_HEIGHT,
        font_size=STACK_TITLE_FONT_SIZE,
        badge_d=0.32,
    )
    photo_h = w * STACK_PHOTO_ASPECT
    photo_y = y + h - 0.20 - photo_h
    _add_step_photo(slide, step, x, photo_y, w, photo_h)

    bullets = split_instructions(step)
    bullet_top = y + STACK_TITLE_BOX_HEIGHT + 0.08
    _add_bullets(
        slide,
        idx,
        bullets,
        _bullet_heights(bullets, w - 0.26, 0.24),
        x,
        bullet_top,
        w - 0.26,
        photo_y - 0.12 - bullet_top,
    )
    add_rule(slide, x, y + h - 0.05, w, C["line"], 0.55)


def photo_fraction(steps_on_page: list[dict[str, Any]]) -> float:
    """Share of a side-by-side block the photograph takes, set per page.

    A fixed share left short steps with a narrow strip of text beside an
    equally narrow photograph and the page's spare height as air. When every
    step on a page is short the photograph widens instead, which also lets it
    grow taller before reaching its portrait cap. One share per page keeps the
    photographs in a row the same width.
    """
    longest = max(
        (sum(len(b) for b in split_instructions(step)) for step in steps_on_page), default=0
    )
    for limit, fraction in PHOTO_FRACTION_BY_TEXT:
        if longest < limit:
            return fraction
    return STEP_IMAGE_WIDTH_FRACTION


def _side_columns(count: int, wide: bool = False) -> tuple[tuple[float, float], ...]:
    # One step alone on a page uses the full width rather than leaving an empty
    # second column beside it, which reads as a rendering fault.
    return ((0.34, 9.22),) if count == 1 or wide else ((0.34, 4.45), (5.18, 4.38))


class SidePlan(NamedTuple):
    """How one page of side-by-side steps is set out."""

    columns: tuple[tuple[float, float], ...]
    fraction: float
    rows: list[float]
    caps: list[float]

    @property
    def needed(self) -> float:
        return sum(self.rows) + STEP_BLOCK_GAP * max(0, len(self.rows) - 1)


def _measure_side(steps_on_page: list[dict[str, Any]], wide: bool) -> SidePlan:
    columns = _side_columns(len(steps_on_page), wide)
    fraction = WIDE_PHOTO_FRACTION if len(columns) == 1 else photo_fraction(steps_on_page)
    rows, caps = [], []
    for start in range(0, len(steps_on_page), len(columns)):
        row = list(zip(steps_on_page[start : start + len(columns)], columns, strict=False))
        rows.append(max(step_block_height(step, width, fraction) for step, (_, width) in row))
        caps.append(
            max(
                step_block_height(step, width, fraction, _tallest_aspect(width))
                for step, (_, width) in row
            )
        )
    return SidePlan(columns, fraction, rows, caps)


def side_plan(steps_on_page: list[dict[str, Any]], room: float) -> SidePlan:
    """Full-width rows when the page has the room for them, else two columns.

    Two steps side by side on a page with room for twice their height left
    half the page empty under them. Set one to a row across the page, each
    gets a wider column of text and a landscape photograph, and the page is
    filled with the recipe rather than with notes or white.
    """
    if 1 < len(steps_on_page) <= WIDE_MAX_STEPS:
        wide = _measure_side(steps_on_page, wide=True)
        if wide.needed <= room:
            return wide
    return _measure_side(steps_on_page, wide=False)


def place_step_blocks(
    steps_on_page: list[dict[str, Any]], *, top: float, bottom: float
) -> list[tuple[float, float, float, float]]:
    """Lay out one page of side-by-side steps, returning each block as (x, y, w, h).

    Separate from drawing so the arithmetic can be checked without building a
    slide: what a block ends up measuring is the whole of this layout's
    behaviour, and it is not observable from the rendered shapes when a step
    has no photograph to place.
    """
    plan = side_plan(steps_on_page, bottom - top)
    # Steps sit in rows that share a top and a bottom, so the rules under two
    # neighbouring steps line up. Independent columns let them drift apart and
    # the grid read as broken.
    # Spare height goes to the photographs, and only until they reach their
    # tallest proportion: stretching the block past that left a band of air
    # under every step rather than more of the dish.
    scale = min(STEP_BLOCK_MAX_SCALE, (bottom - top) / plan.needed) if plan.needed else 1.0
    scale = max(1.0, scale)
    placed: list[tuple[float, float, float, float]] = []
    y = top
    for index, (natural, cap) in enumerate(zip(plan.rows, plan.caps, strict=True)):
        height = min(max(natural, min(natural * scale, cap)), bottom - y)
        for x, width in plan.columns[: len(steps_on_page) - index * len(plan.columns)]:
            placed.append((x, y, width, height))
        y += height + STEP_BLOCK_GAP
    return placed


def _stacked_rows(steps_on_page: list[dict[str, Any]]) -> list[float]:
    width = stacked_column_width()
    heights = [stacked_block_height(step, width) for step in steps_on_page]
    return [
        max(heights[start : start + STACK_COLUMNS])
        for start in range(0, len(heights), STACK_COLUMNS)
    ]


def _stacked_needed(steps_on_page: list[dict[str, Any]]) -> float:
    rows = _stacked_rows(steps_on_page)
    return sum(rows) + STEP_BLOCK_GAP * max(0, len(rows) - 1)


def place_stacked_blocks(
    steps_on_page: list[dict[str, Any]], *, top: float, bottom: float
) -> list[tuple[float, float, float, float]]:
    """Lay out compact steps in rows of three, each row as tall as its tallest."""
    rows = _stacked_rows(steps_on_page)
    needed = _stacked_needed(steps_on_page)
    scale = max(1.0, min(STACK_BLOCK_MAX_SCALE, (bottom - top) / needed)) if needed else 1.0
    width = stacked_column_width()
    placed = []
    y = top
    for row, height in enumerate(rows):
        for column in range(STACK_COLUMNS):
            if row * STACK_COLUMNS + column < len(steps_on_page):
                placed.append((CONTENT_X + column * (width + STACK_GAP), y, width, height * scale))
        y += height * scale + STEP_BLOCK_GAP
    return placed


# ---------------------------------------------------------------------------
# Variations: a short list, or every customization written out in full
# ---------------------------------------------------------------------------


def _line_height(size: float) -> float:
    return size * LINE_HEIGHT_RATIO / POINTS_PER_INCH


def _text_height(text: str, width: float, size: float, *, bold: bool = False) -> float:
    ratio = CHAR_WIDTH_RATIO * (BOLD_WIDTH_FACTOR if bold else 1.0)
    return wrapped_line_count(text, width, size, ratio) * _line_height(size)


def _inline_columns(recipe) -> tuple[list[dict[str, Any]], float, bool]:
    """The options, the width of each one's column, and whether a sketch fits."""
    options = list(recipe.get("customizations") or [])
    count = max(1, min(3, len(options)))
    sketch = count < 3
    area = CONTENT_W - 0.52 - (2.10 if sketch else 0.0)
    return options, (area - (count - 1) * INLINE_COLUMN_GAP) / count, sketch


def _swap_text(option: dict[str, Any]) -> str:
    items = ", ".join(
        f"{clean(ing.get('quantity'))} {clean(ing.get('item'))}".strip()
        for ing in option.get("ingredients", [])
    )
    replaces = ", ".join(option.get("replaces", []))
    if replaces and items:
        return f"Instead of {replaces}: {items}"
    return f"Add {items}" if items else ""


def _draw_option(slide, option: dict[str, Any], x: float, y: float, w: float) -> float:
    """Draw one customization in its column; with no slide, only measure it."""
    cy = y
    name = clean(option.get("name")).upper()
    height = _text_height(name, w, 10, bold=True)
    if slide is not None:
        add_text(
            slide,
            name,
            x,
            cy,
            w,
            height + 0.02,
            font_size=10,
            bold=True,
            color=C["dark_blue"],
            valign="top",
            fit=False,
        )
    cy += height + 0.04
    swap = _swap_text(option)
    if swap:
        height = _text_height(swap, w, 8.5)
        if slide is not None:
            add_text(
                slide,
                swap,
                x,
                cy,
                w,
                height + 0.02,
                font_size=8.5,
                italic=True,
                color=C["muted"],
                valign="top",
                fit=False,
            )
        cy += height + 0.07
    for change in option.get("steps", []):
        if slide is not None:
            add_text(
                slide,
                f"STEP {change['step']}",
                x,
                cy,
                w,
                0.15,
                font_size=7.5,
                bold=True,
                color=C["dark_blue"],
                fit=False,
            )
        cy += 0.18
        for instruction in change.get("instructions", []):
            height = _text_height(instruction, w - 0.22, INLINE_FONT_SIZE)
            if slide is not None:
                add_checkbox(slide, x, cy + 0.025, 0.10)
                add_text(
                    slide,
                    instruction,
                    x + 0.20,
                    cy - 0.01,
                    w - 0.22,
                    height + 0.04,
                    font_size=INLINE_FONT_SIZE,
                    valign="top",
                    margin=0.0,
                    fit=False,
                )
            cy += height + 0.05
        cy += 0.04
    return cy - y


def _inline_panel_height(recipe) -> float:
    options, width, _ = _inline_columns(recipe)
    return 0.24 + max(_draw_option(None, option, 0.0, 0.0, width) for option in options) + 0.14


def customizations_inline(recipe) -> bool:
    """Whether every customization fits written out in the variations panel.

    Written out beside the steps they change, a cook reads them without
    turning to a page of their own; only a set too long for the panel is
    moved to Customized Steps pages.
    """
    return bool(recipe.get("customizations")) and (
        _inline_panel_height(recipe) <= VARIATIONS_INLINE_MAX_H
    )


LIST_TEXT_W = 6.30
LIST_FONT_SIZE = 10.7


def _variation_lines(recipe) -> list[str]:
    """The short list: one line per option, or the legacy variation strings."""
    return (
        [
            option["name"]
            + ": see Customized Steps "
            + ", ".join(str(change["step"]) for change in option["steps"])
            for option in recipe.get("customizations", [])
        ]
        or list(recipe.get("variations") or [])
    )[:3]


def _variation_line_heights(lines: list[str]) -> list[float]:
    # Each entry wraps to the height it needs. Cutting at a fixed length ended
    # a gluten-free swap mid-sentence, before the instruction that mattered.
    return [_text_height(line, LIST_TEXT_W - 0.14, LIST_FONT_SIZE) + 0.04 for line in lines]


def variations_panel_height(recipe) -> float:
    if customizations_inline(recipe):
        return _inline_panel_height(recipe)
    heights = _variation_line_heights(_variation_lines(recipe))
    return max(VARIATIONS_PANEL_H, 0.30 + sum(h + 0.09 for h in heights) + 0.18)


def closing_height(recipe) -> float:
    """Height of the closing panels: the variations with their tab, then the banner."""
    return VARIATIONS_PANEL_OFFSET + variations_panel_height(recipe) + CLOSING_GAP + BANNER_H


def add_variations_panel(slide, recipe, y):
    # The heading sits above the panel as a tab rather than inside it, so the
    # panel starts lower and carries only the list.
    panel_y = y + VARIATIONS_PANEL_OFFSET
    panel_h = variations_panel_height(recipe)
    panel = add_box(slide, 0.34, panel_y, 9.22, panel_h, C["pale"], C["pale"], radius=True)
    add_soft_shadow(panel)
    # A filled tab over the panel's corner, with the heading reversed out of it.
    add_box(
        slide,
        0.34,
        y,
        VARIATIONS_TAB_W,
        VARIATIONS_TAB_H,
        VARIATIONS_TAB_FILL,
        VARIATIONS_TAB_FILL,
        radius=True,
    )
    add_text(
        slide,
        clean(recipe.get("variations_title"), "Simple Variations"),
        0.34,
        y,
        VARIATIONS_TAB_W,
        VARIATIONS_TAB_H,
        font_face=HEAD_FONT,
        font_size=22,
        color=C["white"],
        bold=True,
        align="center",
        margin=0.01,
    )
    sketch_path = recipe.get("variations_image_path") or recipe.get("decorative_image_path")

    if customizations_inline(recipe):
        options, width, sketch = _inline_columns(recipe)
        for index, option in enumerate(options):
            x = 0.60 + index * (width + INLINE_COLUMN_GAP)
            if index:
                add_vrule(
                    slide,
                    x - INLINE_COLUMN_GAP / 2,
                    panel_y + 0.22,
                    panel_h - 0.40,
                    C["border"],
                    0.6,
                )
            _draw_option(slide, option, x, panel_y + 0.24, width)
        if sketch:
            add_ink_image(
                slide, sketch_path, 9.56 - 2.10, panel_y + (panel_h - 1.09) / 2, 2.00, 1.09
            )
        return

    variations = _variation_lines(recipe)
    if not variations:
        # The panel is part of the fixed layout, so an empty one prints as a
        # blank cream box under its heading.
        raise ContentTooLong(
            field="customizations",
            detail="The Simple Variations panel has nothing to list.",
            suggestion=(
                "Add two or three customizations, each with name, replaces, ingredients "
                "and the steps it changes."
            ),
        )
    by = panel_y + 0.30
    for variation, height in zip(variations, _variation_line_heights(variations), strict=True):
        add_checkbox(slide, 0.60, by + 0.035, 0.12)
        add_text(
            slide,
            variation,
            0.85,
            by,
            LIST_TEXT_W,
            height,
            font_size=LIST_FONT_SIZE,
            valign="top",
            fit=False,
        )
        by += height + 0.09
    add_ink_image(
        slide, sketch_path, VARIATIONS_SKETCH_X, panel_y + (panel_h - 1.09) / 2, 2.00, 1.09
    )


def add_bottom_banner(slide, recipe, y):
    supplied = clean(recipe.get("bottom_banner_text"))
    text = supplied or "Made at Home\nfor a cozy evening."
    # The box holds two lines; a third runs out of the bottom of the banner.
    if wrapped_line_count(text, BANNER_TEXT_W, BANNER_FONT_SIZE, BANNER_CHAR_WIDTH_RATIO) > 2:
        raise ContentTooLong(
            field="bottom_banner_text",
            detail="The bottom banner text runs past the two lines its box holds.",
            suggestion=(
                "Write exactly two lines of under 22 characters each, separated by a "
                "newline (\\n), with no bullet or separator joining them."
            ),
        )
    banner = add_box(slide, 0.34, y, 9.22, BANNER_H, C["yellow2"], C["yellow2"], radius=True)
    add_soft_shadow(banner)
    # A short navy stroke closes the slogan, as on the reference card.
    add_rule(slide, BANNER_TEXT_X + 0.02, y + 0.94, 0.55, C["dark_blue"], 1.8)
    add_text(
        slide,
        text,
        BANNER_TEXT_X,
        y + 0.17,
        BANNER_TEXT_W,
        0.84,
        font_face=HEAD_FONT,
        font_size=BANNER_FONT_SIZE,
        color=C["dark_blue"],
        bold=True,
        valign="top",
    )
    add_image(
        slide,
        recipe.get("footer_image_path") or recipe.get("hero_image_path"),
        4.02,
        y,
        5.54,
        BANNER_H,
        crop=True,
        placeholder="FOOD BANNER",
    )


# ---------------------------------------------------------------------------
# Pagination
# ---------------------------------------------------------------------------


def _needed(layout: str, steps_on_page: list[dict[str, Any]], room: float) -> float:
    if not steps_on_page:
        return 0.0
    if layout == LAYOUT_STACKED:
        return _stacked_needed(steps_on_page)
    return side_plan(steps_on_page, room).needed


def _fits(recipe, layout, steps, start, count, bottom) -> bool:
    room = bottom - steps_top(cooking_tip_for_steps(recipe, start, count))
    if start == 0:
        room -= ingredient_strip_height(recipe)
    return _needed(layout, steps[start : start + count], room) <= room


def _pack(recipe, steps: list[dict[str, Any]], layout: str) -> list[tuple[int, int]]:
    """Fill pages with as many steps as their measured heights allow.

    A fixed four per page gave a recipe of six short steps a second page
    holding two steps and an empty half page; packing by height fits them on
    one when they are short and still breaks a long recipe sensibly.
    """
    closing_bottom = CLOSING_BOTTOM - closing_height(recipe) - CLOSING_GAP
    # Pages fill a whole row at a time: two steps side by side, or three
    # across. Filling step by step left a fifth step heading a page of its own
    # column and stranded the sixth alone on the next page.
    row = STACK_COLUMNS if layout == LAYOUT_STACKED else 2
    pages: list[tuple[int, int]] = []
    start = 0
    while start < len(steps):
        count = min(row, len(steps) - start)
        while start + count < len(steps):
            more = min(count + row, len(steps) - start)
            if not _fits(recipe, layout, steps, start, more, STEPS_PAGE_BOTTOM):
                break
            count = more
        pages.append((start, count))
        start += count
    if not pages:
        return [(0, 0)]

    # The closing panels go under the last page's steps when there is room;
    # otherwise the fewest whole rows move on with them.
    start, count = pages[-1]
    if _fits(recipe, layout, steps, start, count, closing_bottom):
        return pages
    for keep in range(count - 1, 0, -1):
        if keep % row == 0 and _fits(
            recipe, layout, steps, start + keep, count - keep, closing_bottom
        ):
            return [*pages[:-1], (start, keep), (start + keep, count - keep)]
    return [*pages, (len(steps), 0)]


def paginate_steps(recipe, steps: list[dict[str, Any]]) -> tuple[str, list[tuple[int, int]]]:
    """Choose the layout that needs fewer pages, preferring side-by-side on a tie.

    Side-by-side, with a tall photograph beside each checklist, is the house
    look. The compact layout, three steps across with each photograph under
    its text, is used only when it saves a page.
    """
    side = _pack(recipe, steps, LAYOUT_SIDE)
    stacked = _pack(recipe, steps, LAYOUT_STACKED) if len(steps) > 2 else side
    return (LAYOUT_STACKED, stacked) if len(stacked) < len(side) else (LAYOUT_SIDE, side)


def add_continued_panels(slide, steps_end: float) -> None:
    """Fill what a page of long steps leaves under them.

    Steps too long to share a page with the variations end a page early, and
    that page used to finish in a band of bare white. A navy band now says the
    method continues, and when there is room above it a ruled notes panel gives
    the cook somewhere to write, so the page reads as designed rather than cut.
    """
    band_y = CONTINUED_BAND_BOTTOM - CONTINUED_BAND_H
    if band_y - steps_end < CONTINUED_MIN_SPARE:
        return
    notes_top = steps_end + CLOSING_GAP + 0.06
    notes_h = band_y - CLOSING_GAP - notes_top
    if notes_h >= NOTES_MIN_H:
        panel = add_box(slide, 0.34, notes_top, 9.22, notes_h, C["pale"], C["pale"], radius=True)
        add_soft_shadow(panel)
        add_text(
            slide,
            "Kitchen Notes",
            0.60,
            notes_top + 0.16,
            4.0,
            0.36,
            font_face=HEAD_FONT,
            font_size=18,
            color=C["dark_blue"],
            bold=True,
        )
        line_y = notes_top + 0.86
        while line_y < notes_top + notes_h - 0.22:
            add_rule(slide, 0.60, line_y, 8.70, C["border"], 0.6)
            line_y += NOTES_LINE_GAP
    else:
        # Too little room for notes: the band follows the steps directly.
        band_y = steps_end + CLOSING_GAP + 0.06
    band = add_box(
        slide, 0.34, band_y, 9.22, CONTINUED_BAND_H, C["dark_blue"], C["dark_blue"], radius=True
    )
    add_soft_shadow(band)
    add_text(
        slide,
        "Steps continue on the next page  →",
        0.34,
        band_y,
        9.22,
        CONTINUED_BAND_H,
        font_face=HEAD_FONT,
        font_size=16,
        color=C["white"],
        bold=True,
        align="center",
        fit=False,
    )


def add_steps_slide(prs, recipe, layout, start, count, *, is_last, page_number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    steps = list(recipe.get("steps") or [])
    steps_on_page = steps[start : start + count]
    tip = cooking_tip_for_steps(recipe, start, count)
    top = steps_top(tip)
    # The rail's overflow opens the first steps page rather than taking a page.
    strip_h = ingredient_strip_height(recipe) if start == 0 else 0.0
    strip_y = top - 0.06
    top += strip_h

    closing_top = CLOSING_BOTTOM - closing_height(recipe)
    # A page carrying only the closing panels starts them at the top.
    if is_last and not steps_on_page:
        closing_top = top
    bottom = closing_top - CLOSING_GAP if is_last else STEPS_PAGE_BOTTOM

    columns = 2
    if layout == LAYOUT_STACKED:
        placed = place_stacked_blocks(steps_on_page, top=top, bottom=bottom)
        draw = add_stacked_block
        columns = STACK_COLUMNS
    elif steps_on_page:
        plan = side_plan(steps_on_page, bottom - top)
        columns = len(plan.columns)
        placed = place_step_blocks(steps_on_page, top=top, bottom=bottom)
        draw = functools.partial(add_step_block, fraction=plan.fraction)
        if columns > 1 and len(steps_on_page) > 1:
            rule_bottom = max(y + h for _, y, _, h in placed)
            add_vrule(slide, 4.98, top - 0.10, rule_bottom - top + 0.10, C["line"], 0.75)
    else:
        placed, draw = [], add_step_block
    add_top_tip(slide, tip, split=columns > 1 and not strip_h)
    if strip_h:
        add_ingredient_strip(slide, recipe, strip_y)

    # The closing panels follow the steps rather than sitting at the foot of
    # the page, so a short last page ends in white below the banner instead
    # of a gap between the steps and the variations.
    if is_last and placed:
        closing_top = min(closing_top, max(y + h for _, y, _, h in placed) + CLOSING_GAP + 0.06)

    inline = customizations_inline(recipe)
    for i, step in enumerate(steps_on_page):
        affected = any(
            change["step"] == start + i + 1
            for option in recipe.get("customizations", [])
            for change in option["steps"]
        )
        if affected:
            step = {**step, "title": clean(step.get("title")) + " *"}
        draw(slide, step, start + i, *placed[i])
    if recipe.get("customizations") and steps_on_page:
        # A footnote under the steps it marks, where a reader looks for one;
        # at the head of the page it sat against the trim with nothing above.
        add_text(
            slide,
            "* Changes for each variation are under Simple Variations."
            if inline
            else "* See Customized Steps for alternatives.",
            0.34,
            max(y + h for _, y, _, h in placed) + 0.02,
            8.8,
            0.16,
            font_size=8,
            italic=True,
            color=C["dark_blue"],
        )

    if is_last:
        # Anchored above the footer, so the banner sits in the same place on
        # every card and the steps take whatever height is left above it.
        add_variations_panel(slide, recipe, closing_top)
        add_bottom_banner(
            slide,
            recipe,
            closing_top + VARIATIONS_PANEL_OFFSET + variations_panel_height(recipe) + CLOSING_GAP,
        )
        add_footer(slide, recipe)
    else:
        if placed:
            add_continued_panels(slide, max(y + h for _, y, _, h in placed))
        # The same rule as the closing footer, so a middle page ends finished.
        add_rule(slide, 0.34, 12.74, 9.22, C["line"], 0.75)
        add_text(
            slide,
            f"{clean(recipe.get('title'))} — continued",
            0.34,
            12.82,
            3.80,
            0.14,
            font_size=7,
            color=C["muted"],
        )
        add_text(
            slide,
            f"Page {page_number}",
            8.91,
            12.82,
            0.65,
            0.14,
            font_size=7,
            color=C["muted"],
            align="right",
        )
    return slide
