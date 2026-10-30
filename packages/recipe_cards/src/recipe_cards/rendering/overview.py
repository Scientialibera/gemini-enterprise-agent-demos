"""Recipe overview layout helpers."""

from __future__ import annotations

import json
from pathlib import Path

from pptx.oxml import parse_xml
from pptx.util import Inches

from .content import heading_size, title_font
from .drawing import (
    add_box,
    add_bullet_list,
    add_image,
    add_ink_image,
    add_rule,
    add_soft_shadow,
    add_text,
    add_vrule,
    bullet_list_height,
    clean,
    gradient_fill,
    limit_text,
    no_shadow,
)
from .footer import add_footer
from .icons import add_clock_icon, add_pot_icon
from .ingredients import add_ingredient_rail
from .theme import (
    BOLD_WIDTH_FACTOR,
    CALLOUT_FONT_SIZES,
    CALLOUT_Y,
    CHAR_WIDTH_RATIO,
    CHEF_NOTE_FONT_SIZE,
    DESCRIPTION_FONT_SIZE,
    DESCRIPTION_H,
    GETTING_STARTED_COL_GAP,
    GETTING_STARTED_COL_W,
    GETTING_STARTED_LINE_H,
    GETTING_STARTED_LIST_FONT_SIZE,
    HEAD_FONT,
    HEADER_PANEL_H,
    HEADER_POT_X,
    HEADER_POT_Y,
    HEADER_TIME_ICON_D,
    HERO_H,
    HERO_W,
    HERO_X,
    LEFT_W,
    MENU_ACCENT_PICTURE_X,
    MENU_BAND_Y,
    MENU_BLURB_H,
    MENU_BLURB_W,
    MENU_PICTURE_H,
    MENU_PICTURE_W,
    MENU_TITLE_FONT_SIZE,
    POINTS_PER_INCH,
    RIGHT_W,
    RIGHT_X,
    SUBTITLE_FONT_SIZE,
    SUBTITLE_H,
    SUBTITLE_W,
    SUBTITLE_Y,
    TITLE_H,
    TITLE_W,
    TITLE_X,
    TITLE_Y,
    C,
)


def add_header_page1(slide, recipe):
    panel = add_box(slide, 0, 0, LEFT_W, HEADER_PANEL_H, C["blue"])
    gradient_fill(panel, C["blue_top"], C["blue_bottom"])
    no_shadow(panel)
    add_rule(slide, 0.30, 0.72, LEFT_W - 0.60, C["blue_soft"], 0.6)
    # The time sits as one centred group rather than hugging the left edge. Its
    # width is measured from the text: a fixed box cut "30 MINUTES" to
    # "30 MINU…" while "45 MINS" fitted.
    time_text = clean(recipe.get("total_time"), "45 MINUTE RECIPE").upper()
    text_w = min(
        LEFT_W - 0.80,
        len(time_text) * 14 * CHAR_WIDTH_RATIO * BOLD_WIDTH_FACTOR / POINTS_PER_INCH + 0.12,
    )
    group_x = (LEFT_W - (HEADER_TIME_ICON_D + 0.10 + text_w)) / 2
    add_clock_icon(slide, group_x + HEADER_TIME_ICON_D / 2, 0.36, HEADER_TIME_ICON_D, C["white"])
    add_text(
        slide,
        time_text,
        group_x + HEADER_TIME_ICON_D + 0.10,
        0.21,
        text_w,
        0.31,
        font_size=14,
        color=C["white"],
        bold=True,
    )
    # Yellow rather than navy: navy on the blue panel all but disappeared, and
    # the yellow mark echoes the servings pill below.
    add_pot_icon(slide, HEADER_POT_X, HEADER_POT_Y, 0.58, C["yellow2"])
    add_text(
        slide,
        recipe.get("title", ""),
        TITLE_X,
        TITLE_Y,
        TITLE_W,
        TITLE_H,
        font_face=HEAD_FONT,
        font_size=title_font(clean(recipe.get("title"))),
        color=C["white"],
        align="center",
        margin=0.01,
    )
    add_text(
        slide,
        recipe.get("subtitle", ""),
        0.35,
        SUBTITLE_Y,
        SUBTITLE_W,
        SUBTITLE_H,
        font_size=SUBTITLE_FONT_SIZE,
        color=C["white"],
        bold=True,
        valign="top",
        align="center",
        margin=0.01,
    )


def _add_getting_started_list(slide, label, items, y, min_height):
    """Draw one labelled list across two columns, returning where it ends.

    A single column made a long pantry list run into the seasonal menu band
    below it, so the items are split down the middle: the block grows sideways
    and its height follows the longer column rather than the whole list.
    """
    add_text(
        slide,
        label,
        RIGHT_X + 0.04,
        y,
        2.30,
        0.16,
        font_size=8.5,
        color=C["dark_blue"],
        bold=True,
    )
    entries = [clean(item) for item in (items or []) if clean(item)]
    list_y = y + 0.21
    if not entries:
        return list_y + min_height

    # Split where the two columns come out closest in height, counting wrapped
    # lines: an even count of items put two one-line tools beside two that
    # wrapped, and the short column ended in a hole.
    def column_height(column):
        return bullet_list_height(column, GETTING_STARTED_COL_W, GETTING_STARTED_LIST_FONT_SIZE)

    split = min(
        range(1, len(entries) + 1),
        key=lambda at: (
            max(column_height(entries[:at]), column_height(entries[at:])),
            -at,
        ),
    )
    columns = [column for column in (entries[:split], entries[split:]) if column]
    # Each column is as tall as its wrapped items, so a long tool name takes
    # its second line instead of being cut to "Colander or…".
    height = max(min_height, *(column_height(column) for column in columns))
    for index, column in enumerate(columns):
        add_bullet_list(
            slide,
            column,
            RIGHT_X + 0.04 + index * (GETTING_STARTED_COL_W + GETTING_STARTED_COL_GAP),
            list_y,
            GETTING_STARTED_COL_W,
            height,
            font_size=GETTING_STARTED_LIST_FONT_SIZE,
        )
    return list_y + height


SEASON_TILE_W = 1.00


def add_overview_right(slide, recipe):
    add_image(
        slide,
        recipe.get("hero_image_path") or recipe.get("heroImagePath"),
        HERO_X,
        0.0,
        HERO_W,
        HERO_H,
        crop=True,
        placeholder="HERO IMAGE",
    )

    # The callout and the Seasonal Menu band run the photograph's full width,
    # so both banners share its left edge and the page's right edge. Inset,
    # their edges sat a little inside the photo's and read as misaligned.
    callout_box = add_box(slide, HERO_X, 7.48, HERO_W, 1.02, C["yellow2"], C["yellow2"])
    add_soft_shadow(callout_box)
    callout = clean(recipe.get("callout_title") or recipe.get("card_title"))
    # The dish name is already the header's headline, so repeating it here, a
    # few inches below, reads as a template fault rather than a second voice.
    # "Lemon Herb Roast Chicken with Potatoes" under "Lemon Herb Roast
    # Chicken" is the same repeat, so a heading that opens with the title
    # counts too.
    if not callout or callout.casefold().startswith(clean(recipe.get("title")).casefold()):
        region = clean(recipe.get("region"))
        callout = f"A Classic {region} Comfort Dish" if region else "Why You'll Love It"
    add_text(
        slide,
        callout,
        RIGHT_X + 0.26,
        CALLOUT_Y,
        RIGHT_W - 0.60,
        0.31,
        font_face=HEAD_FONT,
        # One line: a second ran down over the description below it.
        font_size=heading_size(callout, RIGHT_W - 0.60, CALLOUT_FONT_SIZES),
        color=C["dark_blue"],
        bold=True,
    )
    # Two lines at a readable size rather than one cramped line.
    add_text(
        slide,
        limit_text(recipe.get("description"), 210),
        RIGHT_X + 0.26,
        8.06,
        RIGHT_W - 0.52,
        DESCRIPTION_H,
        font_size=DESCRIPTION_FONT_SIZE,
        valign="top",
    )

    add_text(
        slide,
        "Getting Started",
        RIGHT_X + 0.04,
        8.78,
        2.68,
        0.34,
        font_face=HEAD_FONT,
        font_size=22,
        color=C["dark_blue"],
        bold=True,
    )
    tools_bottom = _add_getting_started_list(
        slide, "COOKING TOOLS", recipe.get("tools"), 9.24, GETTING_STARTED_LINE_H
    )
    # The lists take their own height: a fixed minimum left a band of white
    # between a two-item tools list and the pantry heading under it.
    _add_getting_started_list(
        slide, "FROM YOUR PANTRY", recipe.get("pantry"), tools_bottom + 0.22, GETTING_STARTED_LINE_H
    )

    # Stops above the Seasonal Menu band; at 3.25in it ran into the band's top
    # edge and was hidden only because the band is drawn over it.
    add_vrule(slide, RIGHT_X + 2.84, 8.78, 2.65, C["line"], 1.0)
    chef = recipe.get("chef_note") or {}
    add_text(
        slide,
        "Chef's Note",
        RIGHT_X + 3.12,
        8.78,
        2.80,
        0.34,
        font_face=HEAD_FONT,
        font_size=22,
        color=C["dark_blue"],
        bold=True,
    )
    chef_label = clean(chef.get("title"), "FEATURED INGREDIENT").upper()
    # "CHEF'S NOTE" under the "Chef's Note" heading repeats it word for word.
    if chef_label.replace(chr(0x2019), "'") in ("CHEF'S NOTE", "CHEFS NOTE", "CHEF NOTE"):
        chef_label = "TECHNIQUE & FLAVOR"
    add_text(
        slide,
        chef_label,
        RIGHT_X + 3.12,
        9.24,
        2.80,
        0.16,
        font_size=8.4,
        color=C["dark_blue"],
        bold=True,
    )
    add_text(
        slide,
        limit_text(chef.get("text"), 260),
        RIGHT_X + 3.12,
        9.48,
        2.55,
        1.55,
        font_size=CHEF_NOTE_FONT_SIZE,
        valign="top",
        margin=0.01,
    )
    add_ink_image(
        slide,
        recipe.get("decorative_image_path") or recipe.get("decorativeImagePath"),
        # Ends on the band's right edge: at RIGHT_X + 4.42 it ran 0.3in past
        # the column into the page margin.
        RIGHT_X + RIGHT_W - 0.12 - 1.80,
        10.38,
        1.80,
        1.28,
    )

    banner_y = MENU_BAND_Y
    # Derived from the band rather than repeated as literals, so the picture
    # cannot drift outside it: it used to overhang the right edge by a third of
    # an inch because the two were positioned independently.
    # The same edges as the callout and the photograph above it.
    band_x = HERO_X
    band_w = HERO_W
    band_h = MENU_PICTURE_H
    picture_w = MENU_PICTURE_W
    picture_x = band_x + band_w - picture_w
    band = add_box(slide, band_x, banner_y, band_w, band_h, C["cream"], C["cream"])
    add_soft_shadow(band)
    add_box(slide, band_x, banner_y, SEASON_TILE_W, band_h, C["yellow2"], C["yellow2"])
    add_text(
        slide,
        clean(recipe.get("season")).upper() or "IN\nSEASON",
        band_x,
        banner_y,
        SEASON_TILE_W,
        band_h,
        font_size=10.5,
        color=C["dark_blue"],
        bold=True,
        align="center",
    )
    add_text(
        slide,
        "Seasonal Menu",
        band_x + SEASON_TILE_W + 0.20,
        banner_y + 0.07,
        1.70,
        0.16,
        font_size=MENU_TITLE_FONT_SIZE,
        color=C["dark_blue"],
        bold=True,
    )
    add_text(
        slide,
        limit_text(recipe.get("seasonal_blurb") or recipe.get("description"), 120),
        band_x + SEASON_TILE_W + 0.20,
        banner_y + 0.25,
        MENU_BLURB_W,
        MENU_BLURB_H,
        font_size=7.7,
        color=C["dark_blue"],
        valign="top",
    )
    # The photograph fills the band's full height at its right end.
    add_image(
        slide,
        recipe.get("footer_image_path") or recipe.get("hero_image_path"),
        picture_x,
        banner_y,
        picture_w,
        MENU_PICTURE_H,
        crop=True,
        placeholder="FOOD",
    )
    # Preserve the user's freeform geometry, gradient and stacking order. The
    # shapes were drawn for a photograph starting at MENU_ACCENT_PICTURE_X, so
    # they move with the photograph: left at their stored position, the blue
    # wedge stopped short of the band's edge once the band was realigned.
    shift = int(Inches(picture_x - MENU_ACCENT_PICTURE_X))
    for markup in json.loads(Path(__file__).with_name("banner_shapes.json").read_text()):
        element = parse_xml(markup)
        offset = element.xpath("./p:spPr/a:xfrm/a:off")[0]
        offset.set("x", str(int(offset.get("x")) + shift))
        identity = element.xpath("./p:nvSpPr/p:cNvPr")[0]
        identity.set("id", str(slide.shapes._next_shape_id))
        identity.set("name", "Menu banner accent")
        slide.shapes._spTree.insert_element_before(element, "p:extLst")


def add_overview_slide(prs, recipe):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_page1(slide, recipe)
    add_ingredient_rail(slide, recipe)
    add_overview_right(slide, recipe)
    add_footer(slide, recipe)
    return slide
